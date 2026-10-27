"""Importación desde Excel (.xlsx): catálogos, proveedores y productos.
Cada importación permite elegir: solo agregar filas nuevas, o además actualizar las que ya existen."""
import io, math, re, unicodedata
from datetime import datetime
from flask import Blueprint, request, jsonify, g as flask_g
from openpyxl import load_workbook
from extensions import db
from models import ImportLog, Product, ProductType, ProductGroup, Warehouse, Stock, StockMovement, Supplier
from utils import err, login_required

bp = Blueprint("imports", __name__)
MAX_ROWS = 20000


class ImportErr(Exception):
    pass


# ---------- utilidades ----------
def _norm(s):
    """'Depos.Estand' -> 'depos estand' (sin tildes, mayúsculas ni signos)."""
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def _clean(v):
    if isinstance(v, str):
        v = v.strip()
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return None if v in (None, "") else v


def _txt(v, n):
    v = _clean(v)
    return None if v is None else str(v)[:n]


def _code(v, width=0):
    """Código en mayúsculas; si es numérico se rellena con ceros (1 -> '01')."""
    v = _clean(v)
    if v is None:
        return ""
    s = str(v).strip().upper()
    return s.zfill(width) if width and s.isdigit() else s


def _bool(v, default):
    return default if _clean(v) is None else _norm(v) in ("si", "s", "yes", "y", "1", "true", "activo")


def _qty(v):
    v = _clean(v)
    if v is None:
        return 0.0
    try:
        x = float(v.replace(",", ".")) if isinstance(v, str) else float(v)
    except ValueError:
        return None
    return round(x, 3) if math.isfinite(x) and x >= 0 else None


def _workbook():
    f = request.files.get("file")
    if not f or not f.filename:
        raise ImportErr("Elegí un archivo .xlsx")
    if not f.filename.lower().endswith(".xlsx"):
        raise ImportErr("El archivo debe ser un Excel .xlsx")
    try:
        return load_workbook(io.BytesIO(f.read()), read_only=True, data_only=True)
    except Exception:
        raise ImportErr("No se pudo leer el archivo. Verificá que sea un .xlsx válido.")


def _sheet(wb, wanted):
    if wanted:
        for n in wb.sheetnames:
            if _norm(n) == wanted:
                return wb[n]
        if len(wb.sheetnames) > 1:
            raise ImportErr(f"No se encontró la hoja {wanted.upper()} en el archivo.")
    return wb[wb.sheetnames[0]]


def _table(ws, required):
    """Busca la fila de encabezados (en las primeras 15) y devuelve ({nombre_normalizado: columna}, filas)."""
    need = [_norm(h) for h in required]
    cols, rows, found = {}, [], False
    for i, r in enumerate(ws.iter_rows(values_only=True), start=1):
        if i > MAX_ROWS + 20:
            raise ImportErr(f"El archivo supera el máximo de {MAX_ROWS} filas.")
        if not found:
            if i > 15:
                break
            names = [_norm(c) for c in r]
            if all(n in names for n in need):
                found = True
                for j, n in enumerate(names):
                    if n and n not in cols:
                        cols[n] = j
            continue
        if any(c not in (None, "") for c in r):
            rows.append((i, r))
    if not found:
        raise ImportErr("No se encontró la fila de encabezados. El archivo debe tener las columnas: " + ", ".join(required) + ".")
    return cols, rows


def _get(r, cols, name):
    j = cols.get(name)
    return _clean(r[j]) if j is not None and j < len(r) else None


def _resp(commit, summary, errors, warnings, importable, changes=()):
    return jsonify(commit=commit, summary=summary, importable=importable,
                   errors=errors[:200], error_count=len(errors),
                   warnings=warnings[:100], warning_count=len(warnings),
                   changes=list(changes)[:100], change_count=len(changes))


def _save(kind=None, n_rows=0, n_cols=0, applied=0):
    try:
        if kind:   # deja registro de qué se importó (lo usa la sección Tablas)
            f = request.files.get("file")
            db.session.add(ImportLog(kind=kind, filename=((f.filename if f else "") or "")[:120], rows=n_rows,
                                     source_columns=n_cols, applied=applied, user_id=flask_g.user.id))
        db.session.commit()
        return None
    except Exception:
        db.session.rollback()
        return err("No se pudo guardar la importación. No se aplicó ningún cambio.", 500)


LABELS = {"description": "Descripción", "unit": "Unidad", "type": "Tipo", "group": "Grupo",
          "reorder_point": "Punto de pedido", "standard_supplier": "Proveedor", "lot_required": "Lote requerido",
          "active": "Activo", "name": "Razón social", "trade_name": "Nombre de fantasía", "tax_id": "CUIT/CUIL",
          "email": "E-mail", "phone": "Teléfono", "contact": "Contacto", "blocked": "Bloqueado"}


def _fmt(v):
    if v is None or v == "":
        return "(vacío)"
    if isinstance(v, bool):
        return "Sí" if v else "No"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _diff(obj, changes):
    """['Descripción: A → B', ...]: qué cambiaría en una fila existente."""
    return [f"{LABELS.get(k, k)}: {_fmt(getattr(obj, k))} → {_fmt(v)}" for k, v in changes.items()]


def _chips(label, n_new, n_upd, n_same, update):
    if update:
        return [[f"{label} nuevos", n_new], [f"{label} actualizados", n_upd], [f"{label} sin cambios", n_same]]
    return [[f"{label} nuevos", n_new], [f"{label} ya existentes (se omiten)", n_same]]


# ---------- validación de una fila de producto ----------
def _check_product(r, cols, ctx):
    g = lambda k: _get(r, cols, k)
    code, desc, unit = _code(g("codigo")), _txt(g("descripcion"), 255), _txt(g("unidad"), 20)
    tipo, grupo, dep = _code(g("tipo")), _code(g("grupo"), 4), _code(g("depos estand"), 2)
    msgs, warns = [], []
    if not code:
        msgs.append("Código vacío")
    elif len(code) > 40:
        msgs.append("Código demasiado largo (máx. 40)")
    if not desc:
        msgs.append("Descripción vacía")
    if not unit:
        msgs.append("Unidad vacía")
    if tipo not in ctx["types"]:
        msgs.append(f"Tipo '{tipo}' no está en el catálogo")
    if grupo not in ctx["groups"]:
        msgs.append(f"Grupo '{grupo}' no está en el catálogo")
    if dep not in ctx["whs"]:
        msgs.append(f"Depósito '{dep}' no está en el catálogo")
    rp = _qty(g("punto pedido"))
    if rp is None:
        msgs.append("Punto de pedido inválido")
    sup = _code(g("prove estand")) or None
    if sup and sup not in ctx["sups"]:
        warns.append(f"Proveedor '{sup}' no está cargado (importá primero los proveedores)")
    if msgs:
        return None, msgs, warns
    # campos opcionales que realmente vinieron en el archivo: al actualizar solo se tocan esos
    present = {k for k, col in (("reorder_point", "punto pedido"), ("standard_supplier", "prove estand"),
                                ("lot_required", "seguimiento"), ("active", "ativo")) if g(col) is not None}
    return dict(present=present, code=code, description=desc, unit=unit, type=tipo, group=grupo, dep=dep,
                reorder_point=rp, standard_supplier=sup, lot_required="lote" in _norm(g("seguimiento")),
                active=_bool(g("ativo"), True)), [], warns


# ---------- endpoints ----------
@bp.post("/api/import/catalogs")
@login_required
def import_catalogs():
    """Hoja con secciones DEPOSITOS / TIPOS / GRUPOS (código + descripción). Nunca borra."""
    commit = request.form.get("commit") == "1"
    update_existing = request.form.get("update_existing") == "1"
    sections, sec = {"depositos": [], "tipos": [], "grupos": []}, None
    try:
        ws = _sheet(_workbook(), None)
        for i, r in enumerate(ws.iter_rows(values_only=True), start=1):
            if i > MAX_ROWS:
                raise ImportErr(f"El archivo supera el máximo de {MAX_ROWS} filas.")
            a = _clean(r[0]) if len(r) > 0 else None
            b = _clean(r[1]) if len(r) > 1 else None
            key = _norm(a)
            if key in sections and b is None:
                sec = key
            elif sec and a is not None and b is not None and key not in ("codigo", "clave", "cod grupo"):
                sections[sec].append((i, a, b))
    except ImportErr as e:
        return err(str(e))
    if not any(sections.values()):
        return err("No se encontraron las secciones DEPOSITOS, TIPOS y GRUPOS en el archivo.")

    errors, summary, changes = [], [], []

    def simple(model, rows, width, label):
        cur = {m.code: m for m in model.query}
        n_new = n_upd = n_same = 0
        for i, a, b in rows:
            code, desc = _code(a, width)[:4], str(b)[:80]
            m = cur.get(code)
            if not code:
                errors.append({"row": i, "code": "", "messages": [f"{label}: código vacío"]})
            elif m is None:
                n_new += 1
                cur[code] = model(code=code, description=desc)
                if commit:
                    db.session.add(cur[code])
            elif update_existing and m.description != desc:
                n_upd += 1
                changes.append({"row": i, "code": code, "messages": [f"{label}: {m.description} → {desc}"]})
                if commit:
                    m.description = desc
            else:
                n_same += 1
        summary.extend(_chips(label, n_new, n_upd, n_same, update_existing))
        return n_new + n_upd

    cur = {w.code: w for w in Warehouse.query}
    used = {w.name for w in cur.values()}
    n_new = n_upd = n_same = 0
    for i, a, b in sections["depositos"]:
        code, name = _code(a, 2)[:4], str(b)[:80]
        w = cur.get(code)
        if not code:
            errors.append({"row": i, "code": "", "messages": ["Depósitos: código vacío"]})
        elif w is None:
            if name in used:
                errors.append({"row": i, "code": code, "messages": [f"Ya existe otro depósito llamado '{name}'"]})
                continue
            n_new += 1; used.add(name)
            cur[code] = Warehouse(code=code, name=name)
            if commit:
                db.session.add(cur[code])
        elif update_existing and w.name != name:
            if name in used:
                errors.append({"row": i, "code": code, "messages": [f"Ya existe otro depósito llamado '{name}'"]})
                continue
            n_upd += 1
            changes.append({"row": i, "code": code, "messages": [f"Depósitos: {w.name} → {name}"]})
            used.discard(w.name); used.add(name)
            if commit:
                w.name = name
        else:
            n_same += 1
    summary.extend(_chips("Depósitos", n_new, n_upd, n_same, update_existing))
    changed = n_new + n_upd + simple(ProductType, sections["tipos"], 0, "Tipos") + simple(ProductGroup, sections["grupos"], 4, "Grupos")
    if commit and changed and (resp := _save("catalogs", sum(len(v) for v in sections.values()), 2, changed)):
        return resp
    return _resp(commit, summary, errors, [], changed > 0, changes)


@bp.post("/api/import/suppliers")
@login_required
def import_suppliers():
    """Hoja PROVEEDORES. Al actualizar, las celdas vacías del Excel NO borran lo ya cargado."""
    commit = request.form.get("commit") == "1"
    update_existing = request.form.get("update_existing") == "1"
    try:
        cols, rows = _table(_sheet(_workbook(), "proveedores"), ["Codigo", "Razon Social"])
    except ImportErr as e:
        return err(str(e))
    cur = {s.code: s for s in Supplier.query}
    seen, errors, changes, n_new, n_upd, n_same = set(), [], [], 0, 0, 0
    for i, r in rows:
        g = lambda k: _get(r, cols, k)
        code, name, msgs = _code(g("codigo")), _txt(g("razon social"), 120), []
        if not code:
            msgs.append("Código vacío")
        elif len(code) > 10:
            msgs.append("Código demasiado largo (máx. 10)")
        if not name:
            msgs.append("Razón social vacía")
        if code in seen:
            msgs.append("Código repetido en el archivo")
        seen.add(code)
        if msgs:
            errors.append({"row": i, "code": code, "messages": msgs})
            continue
        data = dict(name=name)
        for key, col, n in (("trade_name", "n fantasia", 120), ("tax_id", "cuit cuil", 20), ("email", "e mail", 120),
                            ("phone", "telefono", 40), ("contact", "contacto", 120)):
            v = _txt(g(col), n)
            if v is not None:
                data[key] = v
        if g("bloqueado") is not None:
            data["blocked"] = _bool(g("bloqueado"), False)
        s = cur.get(code)
        if s is None:
            n_new += 1
            if commit:
                db.session.add(Supplier(code=code, **data))
        elif update_existing:
            ch = {k: v for k, v in data.items() if getattr(s, k) != v}
            if ch:
                n_upd += 1
                changes.append({"row": i, "code": code, "messages": _diff(s, ch)})
                if commit:
                    for k, v in ch.items():
                        setattr(s, k, v)
            else:
                n_same += 1
        else:
            n_same += 1
    if commit and (n_new or n_upd) and (resp := _save("suppliers", len(rows), len(cols), n_new + n_upd)):
        return resp
    summary = _chips("Proveedores", n_new, n_upd, n_same, update_existing) + [["Filas con error (se omiten)", len(errors)]]
    return _resp(commit, summary, errors, [], (n_new + n_upd) > 0, changes)


@bp.post("/api/import/products")
@login_required
def import_products():
    """Hoja PRODUCTOS. Crea productos nuevos con stock 0 en su depósito estándar.
    Al actualizar: nunca toca cantidades ni depósitos, y solo cambia los campos que vienen en el archivo."""
    commit = request.form.get("commit") == "1"
    update_existing = request.form.get("update_existing") == "1"
    try:
        cols, rows = _table(_sheet(_workbook(), "productos"),
                            ["Codigo", "Descripcion", "Tipo", "Unidad", "Grupo", "Depos.Estand"])
    except ImportErr as e:
        return err(str(e))
    ctx = {"types": {t.code for t in ProductType.query}, "groups": {g.code for g in ProductGroup.query},
           "whs": {w.code: w for w in Warehouse.query}, "sups": {s.code for s in Supplier.query}}
    existing = {p.code: p for p in Product.query}
    seen, errors, warnings, changes, n_new, n_upd, n_same = set(), [], [], [], 0, 0, 0
    for i, r in rows:
        data, msgs, warns = _check_product(r, cols, ctx)
        code = data["code"] if data else _code(_get(r, cols, "codigo"))
        if data and code in seen:
            data, msgs = None, ["Código repetido en el archivo"]
        seen.add(code)
        if warns:
            warnings.append({"row": i, "code": code, "messages": warns})
        if not data:
            errors.append({"row": i, "code": code, "messages": msgs})
            continue
        present, dep = data.pop("present"), data.pop("dep")
        p = existing.get(code)
        if p is None:
            n_new += 1
            if commit:
                p = Product(**data)
                db.session.add(p)
                db.session.flush()
                db.session.add(Stock(product_id=p.id, warehouse_id=ctx["whs"][dep].id, quantity=0))
        elif update_existing:
            new = {k: data[k] for k in ("description", "unit", "type", "group")}
            new.update({k: data[k] for k in present})
            ch = {k: v for k, v in new.items() if getattr(p, k) != v}
            if ch:
                n_upd += 1
                changes.append({"row": i, "code": code, "messages": _diff(p, ch)})
                if commit:
                    for k, v in ch.items():
                        setattr(p, k, v)
            else:
                n_same += 1
        else:
            n_same += 1
    if commit and (n_new or n_upd) and (resp := _save("products", len(rows), len(cols), n_new + n_upd)):
        return resp
    summary = _chips("Productos", n_new, n_upd, n_same, update_existing) + [["Filas con error (se omiten)", len(errors)]]
    return _resp(commit, summary, errors, warnings, (n_new + n_upd) > 0, changes)

@bp.post("/api/import/stock")
@login_required
def import_stock():
    """Hoja con Producto, Deposito y Saldo Actual. FIJA la cantidad (saldo) de productos que YA existen, por código.
    Los saldos vienen de TOTVS, así que los movimientos de ajuste quedan como PROCESADOS (no vuelven a la cola de TOTVS)."""
    commit = request.form.get("commit") == "1"
    update_existing = request.form.get("update_existing") == "1"
    try:
        cols, rows = _table(_sheet(_workbook(), None), ["Producto", "Deposito", "Saldo Actual"])
    except ImportErr as e:
        return err(str(e))
    products = {p.code: p for p in Product.query}
    whs = {w.code: w for w in Warehouse.query}
    stocks = {(s.product_id, s.warehouse_id): s for s in Stock.query}
    moved = set(db.session.query(StockMovement.product_id, StockMovement.warehouse_id).distinct().all())
    seen, errors, changes = set(), [], []
    n_new = n_upd = n_same = n_missing = n_bad = 0
    now = datetime.utcnow()
    for i, r in rows:
        code, dep = _code(_get(r, cols, "producto")), _code(_get(r, cols, "deposito"), 2)
        raw = _get(r, cols, "saldo actual")
        qty = None if raw is None else _qty(raw)
        p, w, msgs = products.get(code), whs.get(dep), []
        if not code:
            msgs.append("Producto vacío")
        elif p is None:
            n_missing += 1
            errors.append({"row": i, "code": code, "messages": ["El producto no existe en el sistema (importá primero los productos)"]})
            continue
        elif not p.active:
            msgs.append("El producto está archivado")
        if w is None:
            msgs.append(f"Depósito '{dep}' no está en el catálogo")
        if raw is None:
            msgs.append("Saldo vacío")
        elif qty is None:
            msgs.append("Saldo inválido (debe ser un número mayor o igual a 0)")
        if (code, dep) in seen:
            msgs.append("Producto y depósito repetidos en el archivo")
        seen.add((code, dep))
        if msgs:
            n_bad += 1
            errors.append({"row": i, "code": code, "messages": msgs})
            continue
        k = (p.id, w.id)
        s = stocks.get(k)
        if s is not None and round(s.quantity, 3) == qty:
            n_same += 1
            continue
        has_balance = s is not None and (s.quantity != 0 or k in moved)   # ya tiene un saldo cargado
        if has_balance and not update_existing:
            n_same += 1
            continue
        old = s.quantity if s is not None else 0.0
        if has_balance:
            n_upd += 1
            changes.append({"row": i, "code": code, "messages": [f"Depósito {dep}: {_fmt(old)} → {_fmt(qty)}"]})
        else:
            n_new += 1
        if commit:
            if s is None:
                s = Stock(product_id=p.id, warehouse_id=w.id, quantity=qty)
                db.session.add(s)
                stocks[k] = s
            else:
                s.quantity = qty
            delta = round(qty - old, 3)
            if delta:   # el stock siempre queda explicado por movimientos
                db.session.add(StockMovement(
                    product_id=p.id, warehouse_id=w.id, user_id=flask_g.user.id, quantity=abs(delta),
                    movement_type="IN" if delta > 0 else "OUT", sync_status="PROCESSED", processed_at=now,
                    observation=("Ajuste por importación de saldo: " + f"{old:g} → {qty:g}") if has_balance
                    else "Carga de saldo por importación"))
    if commit and (n_new or n_upd) and (resp := _save("stock", len(rows), len(cols), n_new + n_upd)):       
        return resp
    summary = ([["Saldos a cargar (nuevos)", n_new], ["Saldos actualizados", n_upd], ["Saldos sin cambios", n_same]]
               if update_existing else [["Saldos a cargar (nuevos)", n_new], ["Saldos ya cargados (se omiten)", n_same]])
    summary += [["Productos inexistentes (se omiten)", n_missing], ["Otras filas con error", n_bad]]
    return _resp(commit, summary, errors, [], (n_new + n_upd) > 0, changes)